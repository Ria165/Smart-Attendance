import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def build_excel_attendance_report(report_data):
    """
    Generates a professionally styled Excel workbook containing:
    1. Overview & Summary metrics
    2. Student Attendance Table with conditional formatting for low attendance
    3. Low Attendance Warning & Planning Estimate table
    """
    wb = openpyxl.Workbook()
    
    # ------------------ SHEET 1: ATTENDANCE DETAILS ------------------
    ws = wb.active
    ws.title = "Monthly Attendance"
    ws.views.sheetView[0].showGridLines = True

    # Color definitions (College Navy Theme)
    navy_header_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    sub_header_fill = PatternFill(start_color="2E5B88", end_color="2E5B88", fill_type="solid")
    zebra_fill = PatternFill(start_color="F7F9FC", end_color="F7F9FC", fill_type="solid")
    low_att_fill = PatternFill(start_color="FDE8E8", end_color="FDE8E8", fill_type="solid") # light red
    pass_fill = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid") # light green

    white_bold_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=16, bold=True, color="1B365D")
    sub_title_font = Font(name="Calibri", size=11, italic=True, color="4A5568")
    bold_font = Font(name="Calibri", size=11, bold=True)
    regular_font = Font(name="Calibri", size=11)
    low_font = Font(name="Calibri", size=11, bold=True, color="C53030")
    pass_font = Font(name="Calibri", size=11, bold=True, color="22543D")

    thin_border = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB")
    )

    # Title block
    ws["A1"] = "COLLEGE SMART ATTENDANCE REPORT"
    ws["A1"].font = title_font
    ws["A2"] = f"Period: {report_data['month_name']} {report_data['year']} | Generated: {report_data['generated_at'].strftime('%Y-%m-%d %H:%M UTC')}"
    ws["A2"].font = sub_title_font

    # Metadata cards
    ws["A4"] = "Subject:"
    ws["B4"] = f"{report_data['subject'].subject_code} - {report_data['subject'].subject_name}"
    ws["A5"] = "Class / Section:"
    ws["B5"] = report_data['class'].name
    ws["A6"] = "Faculty In-Charge:"
    ws["B6"] = f"Prof. {report_data['staff'].user.full_name}"

    for r in range(4, 7):
        ws[f"A{r}"].font = bold_font
        ws[f"B{r}"].font = regular_font

    ws["D4"] = "Sessions Held:"
    ws["E4"] = report_data['total_sessions']
    ws["D5"] = "Enrolled Students:"
    ws["E5"] = report_data['total_enrolled']
    ws["D6"] = "Class Average:"
    ws["E6"] = f"{report_data['class_average']}%" if report_data['class_average'] is not None else "N/A"

    for r in range(4, 7):
        ws[f"D{r}"].font = bold_font
        ws[f"E{r}"].font = bold_font

    # Table Header (Row 8)
    headers = [
        "Student ID", "Student Name", "Sessions Held", "Present", "Absent", 
        "Attendance %", "Threshold (75%) Status", "Additional Sessions Needed"
    ]
    for col_idx, text in enumerate(headers, 1):
        cell = ws.cell(row=8, column=col_idx, value=text)
        cell.font = white_bold_font
        cell.fill = navy_header_fill
        cell.alignment = Alignment(horizontal="center" if col_idx > 2 else "left", vertical="center")
        cell.border = thin_border

    # Data Rows
    current_row = 9
    for idx, s in enumerate(report_data["students"]):
        row_cells = [
            ws.cell(row=current_row, column=1, value=s["student_id"]),
            ws.cell(row=current_row, column=2, value=s["student_name"]),
            ws.cell(row=current_row, column=3, value=s["sessions_held"]),
            ws.cell(row=current_row, column=4, value=s["present_count"]),
            ws.cell(row=current_row, column=5, value=s["absent_count"]),
            ws.cell(row=current_row, column=6, value=s["percentage_display"]),
            ws.cell(row=current_row, column=7, value="BELOW THRESHOLD" if s["is_low"] else "ELIGIBLE"),
            ws.cell(row=current_row, column=8, value=s["sessions_needed"] if s["is_low"] else "-")
        ]

        is_even = (idx % 2 == 0)

        for c_idx, c in enumerate(row_cells, 1):
            c.border = thin_border
            c.font = regular_font
            if not is_even:
                c.fill = zebra_fill

            if c_idx in (1, 2):
                c.alignment = Alignment(horizontal="left", vertical="center")
            else:
                c.alignment = Alignment(horizontal="center", vertical="center")

        # Specific styling for status and percentage
        if s["is_low"]:
            row_cells[5].font = low_font
            row_cells[6].font = low_font
            row_cells[6].fill = low_att_fill
            row_cells[7].font = low_font
        else:
            row_cells[6].font = pass_font
            row_cells[6].fill = pass_fill

        current_row += 1

    # ------------------ SHEET 2: LOW ATTENDANCE SUMMARY ------------------
    ws_low = wb.create_sheet(title="Low Attendance Summary")
    ws_low.views.sheetView[0].showGridLines = True

    ws_low["A1"] = "STUDENTS REQUIRING ATTENDANCE REVIEW (< 75%)"
    ws_low["A1"].font = title_font
    ws_low["A2"] = f"Planning estimates to reach {report_data['threshold']:.0f}% minimum attendance requirement"
    ws_low["A2"].font = sub_title_font

    low_headers = ["Student ID", "Student Name", "Held", "Present", "Current %", "Consecutive Sessions Needed (Est.)"]
    for col_idx, text in enumerate(low_headers, 1):
        cell = ws_low.cell(row=4, column=col_idx, value=text)
        cell.font = white_bold_font
        cell.fill = sub_header_fill
        cell.alignment = Alignment(horizontal="center" if col_idx > 2 else "left", vertical="center")
        cell.border = thin_border

    l_row = 5
    if report_data["low_attendance_students"]:
        for s in report_data["low_attendance_students"]:
            cells = [
                ws_low.cell(row=l_row, column=1, value=s["student_id"]),
                ws_low.cell(row=l_row, column=2, value=s["student_name"]),
                ws_low.cell(row=l_row, column=3, value=s["sessions_held"]),
                ws_low.cell(row=l_row, column=4, value=s["present_count"]),
                ws_low.cell(row=l_row, column=5, value=s["percentage_display"]),
                ws_low.cell(row=l_row, column=6, value=f"{s['sessions_needed']} sessions")
            ]
            for c_idx, c in enumerate(cells, 1):
                c.border = thin_border
                c.font = low_font if c_idx in (5, 6) else regular_font
                c.alignment = Alignment(horizontal="center" if c_idx > 2 else "left", vertical="center")
                if c_idx in (5, 6):
                    c.fill = low_att_fill
            l_row += 1
    else:
        ws_low.cell(row=5, column=1, value="All enrolled students meet the attendance eligibility threshold.")
        ws_low.cell(row=5, column=1).font = pass_font

    # Auto-adjust column widths for both sheets
    for sheet in (ws, ws_low):
        for col in sheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                max_len = max(max_len, len(val_str))
            sheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
